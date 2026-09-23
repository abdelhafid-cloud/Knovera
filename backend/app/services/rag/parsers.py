"""Document parsers — extract plain text from supported formats."""

from io import BytesIO

from app.logging_config import get_doc_logger, truncate

logger = get_doc_logger(__name__)


def extract_text(data: bytes, mime_type: str, filename: str = "") -> list[dict]:
    """Return list of {content, page_number, section_title}."""
    name = (filename or "").lower()
    if mime_type == "application/pdf" or name.endswith(".pdf"):
        logger.info("[PIPELINE:EXTRACT] parser=pdf | bytes=%s", len(data))
        return _extract_pdf(data)
    if "wordprocessingml" in mime_type or name.endswith(".docx"):
        logger.info("[PIPELINE:EXTRACT] parser=docx | bytes=%s", len(data))
        return _extract_docx(data)
    if "spreadsheetml" in mime_type or name.endswith(".xlsx"):
        logger.info("[PIPELINE:EXTRACT] parser=xlsx | bytes=%s", len(data))
        return _extract_xlsx(data)
    if mime_type.startswith("text/") or name.endswith(".txt"):
        logger.info("[PIPELINE:EXTRACT] parser=txt | bytes=%s", len(data))
        text = data.decode("utf-8", errors="ignore")
        return [{"content": text, "page_number": 1, "section_title": None}]
    raise ValueError(f"Type de fichier non supporté: {mime_type}")


def _extract_pdf(data: bytes) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    skipped = 0
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append({"content": text, "page_number": i, "section_title": None})
            logger.debug(
                "[PIPELINE:EXTRACT] pdf page=%s | chars=%s | preview=%r",
                i,
                len(text),
                truncate(text),
            )
        else:
            skipped += 1
            logger.warning("[PIPELINE:EXTRACT] pdf page=%s sans texte", i)
    if skipped:
        logger.warning(
            "[PIPELINE:EXTRACT] pdf pages_sans_texte=%s / %s",
            skipped,
            len(reader.pages),
        )
    return pages or [{"content": "", "page_number": 1, "section_title": None}]


def _extract_docx(data: bytes) -> list[dict]:
    from docx import Document as DocxDocument

    doc = DocxDocument(BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    logger.info(
        "[PIPELINE:EXTRACT] docx paragraphs=%s | chars=%s",
        len(paragraphs),
        len(text),
    )
    return [{"content": text, "page_number": 1, "section_title": None}]


def _extract_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    chunks = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = " | ".join(cells).strip(" |")
            if line:
                rows.append(line)
        if rows:
            chunks.append(
                {
                    "content": "\n".join(rows),
                    "page_number": None,
                    "section_title": sheet.title,
                }
            )
            logger.info(
                "[PIPELINE:EXTRACT] xlsx sheet=%s | rows=%s",
                sheet.title,
                len(rows),
            )
    return chunks or [{"content": "", "page_number": None, "section_title": None}]
